"""Editable reports via python-docx — locked (§22 req 14: editable text, not images)."""

from io import BytesIO

from docx import Document
from docx.shared import Pt


def render_docx(dto: dict, evidence_path: str | None = None) -> bytes:
    doc = Document()
    doc.add_heading("MetroScan Compliance Report", 0)
    doc.add_paragraph(dto["disclaimer"])
    m = dto["meta"]
    doc.add_paragraph(f"Scan {m['scan_id']} · {m['filename']} · state {m['processing_state']}")
    doc.add_paragraph(f"Pipeline {m['pipeline_version']} · Rules {m['rule_set_version']} · "
                      f"{m['created_at']}")
    if evidence_path:
        try:
            doc.add_picture(evidence_path, width=Pt(360))
            doc.paragraphs[-1].add_run().add_break()
        except Exception:
            doc.add_paragraph("(original evidence image unavailable)")
    doc.add_heading("Extracted declarations (with OCR confidence)", 1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text, table.rows[0].cells[2].text = \
        "Field", "Value", "OCR confidence"
    for key, f in (dto["fields"] or {}).items():
        row = table.add_row().cells
        row[0].text = key
        row[1].text = f.get("value_raw") or "could not be verified from the provided image"
        row[2].text = f"{float(f.get('ocr_confidence', 0)):.0%}"
    doc.add_heading("Machine findings (pending review)", 1)
    if dto["machine_findings"]:
        for x in dto["machine_findings"]:
            doc.add_paragraph(f"{x['rule_id']} {x['clause']} [{x['machine_status']}] "
                              f"{x['message']}", style="List Bullet")
    else:
        doc.add_paragraph("(none)")
    doc.add_heading("Inspector decisions", 1)
    if dto["human_decisions"]:
        for x in dto["human_decisions"]:
            r = x["review"]
            doc.add_paragraph(f"{x['rule_id']}: {r['reviewer']} → {r['new_status']} "
                              f"at {r['at']} — {r.get('note', '')}", style="List Bullet")
    else:
        doc.add_paragraph("(none)")
    doc.add_heading("Limitations", 1)
    for lim in dto["limitations"]:
        doc.add_paragraph(lim, style="List Bullet")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
